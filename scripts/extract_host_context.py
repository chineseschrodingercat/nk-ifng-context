"""Reconstruct the 28 Fig. 5i mouse observations directly from the published workbook."""
from pathlib import Path
import hashlib, json
import numpy as np
import pandas as pd
from openpyxl import load_workbook

R = Path(__file__).resolve().parents[1]
src = R/'raw/Remsik2025/Remsik2025_main_source.xlsx'
sha = hashlib.sha256(src.read_bytes()).hexdigest()
wb = load_workbook(src, data_only=True, read_only=True)
ws = wb['F5']
rows = []
for col, z, dep, label, ab in [('P',0,0,'eGFP','isotype'),('Q',1,0,'Ifng','isotype'),('R',0,1,'eGFP','anti-asialo-GM1'),('S',1,1,'Ifng','anti-asialo-GM1')]:
    for i in range(5,12):
        cell = f'{col}{i}'
        value = ws[cell].value
        assert isinstance(value,(float,int)) and value > 0
        rows.append(dict(study='Remsik2025', panel='Fig5i', workbook=src.name, sheet='F5', cell=cell, source_sha256=sha, arm_label=label, antibody=ab, ifng_overexpression=z, nk_depletion=dep, mouse_row_id='F5_'+cell, radiance=value, radiance_units='photons s^-1 cm^-2 sr^-1', log10_radiance=np.log10(value), unit='mouse', group_pairing='unpaired; rows across arms are not matched mice', batch_id='not supplied'))
wb.close()
df = pd.DataFrame(rows)
old = pd.read_csv(R/'source_data/Remsik2025_Fig5_mouse_source_cells.csv').query("panel == 'Fig5i'")
join = df.merge(old, on='cell', suffixes=('_new','_old'), validate='one_to_one')
assert len(join)==28 and np.allclose(join.radiance_new, join.radiance_old, rtol=0, atol=0)
assert (join.source_sha256_new==join.source_sha256_old).all()
df.to_csv(R/'analysis/Remsik2025_Fig5i_mouse_values.csv', index=False)
m = df.groupby(['ifng_overexpression','nk_depletion']).log10_radiance.mean()
effect = [float(m[1,d]-m[0,d]) for d in [0,1]]
out = dict(n_mice=28, n_per_arm=7, source_sha256=sha, ifng_effect_log10_isotype=effect[0], ifng_effect_log10_anti_asialo_GM1=effect[1], interaction_depleted_minus_isotype=effect[1]-effect[0], geometric_ratio_isotype=10**effect[0], geometric_ratio_depleted=10**effect[1], inference='descriptive; no pairing or known batch structure; antibody depletion is not exclusive NK attribution')
original = json.loads((R/'source_data/Remsik2025_Fig5i_descriptive_effects.json').read_text())
assert np.isclose(out['interaction_depleted_minus_isotype'],original['interaction_difference_of_log10_effects_depleted_minus_intact'], atol=1e-12)
(R/'analysis/Remsik2025_Fig5i_summary.json').write_text(json.dumps(out,indent=2))
print(json.dumps(out,indent=2))
